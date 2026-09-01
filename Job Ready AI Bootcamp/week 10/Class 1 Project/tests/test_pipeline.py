"""
Tests for the Semantic CV-to-Job Matcher pipeline.

Run with: pytest tests/test_pipeline.py

HIGHLIGHTS: this test file never imports sentence_transformers and never
calls src.embeddings.service.get_model()/embed(). All ranking tests use
small, hand-crafted numpy vectors instead of real embeddings — that's a
deliberate stand-in ("stub") for the embedding service, so these tests run
in milliseconds, need no internet access, and stay deterministic (a real
model's output could theoretically shift across library versions; a
hand-picked vector like [1, 0, 0] never will). See
src/embeddings/service.py's module docstring for the full lazy-loading
rationale this test suite is designed to respect.
"""

from __future__ import annotations

import io
import math

import numpy as np
import pytest
from docx import Document

from src.extraction.text_extractor import (
    ExtractionError,
    extract_text,
    extract_text_from_docx,
    extract_text_from_pdf,
    extract_text_from_plain,
)
from src.matching.ranker import MatchResult, cosine_similarity, rank_cvs


# ─────────────────────────────────────────────────────────────────
# Text extraction: plain text
# ─────────────────────────────────────────────────────────────────

def test_extract_text_from_plain_strips_whitespace():
    result = extract_text_from_plain("   Senior Python Engineer, 5 years experience.   \n")
    assert result == "Senior Python Engineer, 5 years experience."


def test_extract_text_dispatcher_plain():
    result = extract_text("Backend engineer with FastAPI experience.", filetype="plain")
    assert "FastAPI" in result


def test_extract_text_dispatcher_rejects_empty_plain_text():
    with pytest.raises(ExtractionError):
        extract_text("   ", filetype="plain")


def test_extract_text_dispatcher_unsupported_filetype():
    with pytest.raises(ExtractionError):
        extract_text("hello", filetype="rtf")


# ─────────────────────────────────────────────────────────────────
# Text extraction: PDF (built in-memory with pypdf's writer, no fixture
# files on disk — keeps the test fast and self-contained)
# ─────────────────────────────────────────────────────────────────

def _build_minimal_pdf_bytes(text: str) -> bytes:
    """Build a minimal single-page PDF with a real text layer, in memory.

    HIGHLIGHTS: pypdf's PdfWriter can't easily draw arbitrary text onto a
    fresh page on its own (that's normally reportlab's job), but it *can*
    add a page and stamp an existing text-bearing content stream onto it
    via low-level content-stream bytes. Rather than pull in reportlab as
    an extra test-only dependency, we hand-construct a tiny valid PDF
    content stream directly — this keeps the test suite's dependency
    footprint identical to the app's own requirements.txt.
    """
    from pypdf import PdfWriter
    from pypdf.generic import ContentStream, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)

    # A minimal content stream: select a base font, then show text.
    content = f"BT /F1 12 Tf 10 100 Td ({text}) Tj ET".encode("latin-1")
    stream = ContentStream(None, writer)
    stream.set_data(content)
    page[NameObject("/Contents")] = writer._add_object(stream)

    # A page needs a font resource referenced by the content stream above.
    page[NameObject("/Resources")] = writer._add_object(
        DictionaryObject({
            NameObject("/Font"): writer._add_object(
                DictionaryObject({
                    NameObject("/F1"): writer._add_object(
                        DictionaryObject({
                            NameObject("/Type"): NameObject("/Font"),
                            NameObject("/Subtype"): NameObject("/Type1"),
                            NameObject("/BaseFont"): NameObject("/Helvetica"),
                        })
                    )
                })
            )
        })
    )

    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def test_extract_text_from_pdf_reads_text_layer():
    pdf_bytes = _build_minimal_pdf_bytes("Data Scientist")
    result = extract_text_from_pdf(pdf_bytes)
    assert "Data Scientist" in result


def test_extract_text_from_pdf_no_text_layer_raises_gracefully():
    """A PDF with a blank page (no content stream at all) has no text
    layer — this simulates a scanned/image-only PDF. Extraction must
    raise a clear ExtractionError, not crash or silently return "".
    """
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)  # no /Contents at all
    buf = io.BytesIO()
    writer.write(buf)

    with pytest.raises(ExtractionError):
        extract_text_from_pdf(buf.getvalue())


def test_extract_text_from_pdf_invalid_bytes_raises():
    with pytest.raises(ExtractionError):
        extract_text_from_pdf(b"this is not a pdf")


# ─────────────────────────────────────────────────────────────────
# Text extraction: DOCX (built in-memory with python-docx's own
# Document() — no fixture files on disk)
# ─────────────────────────────────────────────────────────────────

def _build_minimal_docx_bytes(paragraphs: list[str]) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extract_text_from_docx_reads_paragraphs():
    docx_bytes = _build_minimal_docx_bytes(
        ["Jane Doe", "Machine Learning Engineer", "5 years of experience with PyTorch."]
    )
    result = extract_text_from_docx(docx_bytes)
    assert "Jane Doe" in result
    assert "PyTorch" in result


def test_extract_text_from_docx_empty_document_raises():
    docx_bytes = _build_minimal_docx_bytes([])
    with pytest.raises(ExtractionError):
        extract_text_from_docx(docx_bytes)


def test_extract_text_dispatcher_docx():
    docx_bytes = _build_minimal_docx_bytes(["Full-stack developer, React and Node.js."])
    result = extract_text(docx_bytes, filetype="docx")
    assert "React" in result


# ─────────────────────────────────────────────────────────────────
# Cosine similarity correctness (hand-computed small vectors)
# ─────────────────────────────────────────────────────────────────

def test_cosine_similarity_identical_vectors_is_one():
    v = np.array([1.0, 2.0, 3.0])
    assert cosine_similarity(v, v) == pytest.approx(1.0)


def test_cosine_similarity_orthogonal_vectors_is_zero():
    a = np.array([1.0, 0.0])
    b = np.array([0.0, 1.0])
    assert cosine_similarity(a, b) == pytest.approx(0.0)


def test_cosine_similarity_opposite_vectors_is_minus_one():
    a = np.array([1.0, 0.0])
    b = np.array([-1.0, 0.0])
    assert cosine_similarity(a, b) == pytest.approx(-1.0)


def test_cosine_similarity_ignores_magnitude():
    """A short vector and a much longer one, same direction, should score
    identically to 1.0 — this is exactly the "long CV vs. short JD" case
    the module docstring explains cosine similarity is chosen to handle.
    """
    a = np.array([1.0, 1.0, 1.0])
    b = a * 50.0  # same direction, 50x the magnitude
    assert cosine_similarity(a, b) == pytest.approx(1.0)


def test_cosine_similarity_zero_vector_returns_zero_not_nan():
    a = np.array([0.0, 0.0, 0.0])
    b = np.array([1.0, 2.0, 3.0])
    result = cosine_similarity(a, b)
    assert result == 0.0
    assert not math.isnan(result)


def test_cosine_similarity_known_angle():
    # 45-degree angle vectors -> cosine similarity should be ~0.7071
    a = np.array([1.0, 0.0])
    b = np.array([1.0, 1.0])
    assert cosine_similarity(a, b) == pytest.approx(math.sqrt(2) / 2, rel=1e-6)


# ─────────────────────────────────────────────────────────────────
# Ranking order correctness (stubbed embeddings — no real model loaded)
# ─────────────────────────────────────────────────────────────────

def test_rank_cvs_orders_best_match_first():
    job_desc = np.array([1.0, 0.0, 0.0])
    cv_embeddings = np.array(
        [
            [0.0, 1.0, 0.0],  # orthogonal -> unrelated, score ~0
            [1.0, 0.0, 0.0],  # identical direction -> perfect match, score 1
            [0.7, 0.7, 0.0],  # partial overlap -> middling score
        ]
    )
    names = ["Unrelated CV", "Perfect Match CV", "Partial Match CV"]

    results = rank_cvs(job_desc, cv_embeddings, names)

    assert [r.name for r in results] == ["Perfect Match CV", "Partial Match CV", "Unrelated CV"]
    assert results[0].rank == 1
    assert results[0].score == pytest.approx(1.0)
    assert results[-1].score == pytest.approx(0.0, abs=1e-9)


def test_rank_cvs_returns_match_result_dataclass_instances():
    job_desc = np.array([1.0, 0.0])
    cv_embeddings = np.array([[1.0, 0.0]])
    results = rank_cvs(job_desc, cv_embeddings, ["Solo CV"])
    assert isinstance(results[0], MatchResult)
    assert results[0].rank == 1


def test_rank_cvs_mismatched_lengths_raises_value_error():
    job_desc = np.array([1.0, 0.0])
    cv_embeddings = np.array([[1.0, 0.0], [0.0, 1.0]])
    with pytest.raises(ValueError):
        rank_cvs(job_desc, cv_embeddings, ["Only One Name"])


def test_rank_cvs_ranks_are_sequential_starting_at_one():
    job_desc = np.array([1.0, 0.0])
    cv_embeddings = np.array([[0.1, 0.9], [0.9, 0.1], [0.5, 0.5]])
    names = ["A", "B", "C"]
    results = rank_cvs(job_desc, cv_embeddings, names)
    assert [r.rank for r in results] == [1, 2, 3]
